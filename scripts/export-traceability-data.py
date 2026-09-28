#!/usr/bin/env python3
"""Dump the Requirements + Traceability Report tabs of docs/KRANE-SRS-traceability.xlsx
to docs/traceability-data.json, so generate-traceability-report.js can read plain JSON
instead of depending on an xlsx-parsing library. Re-run this whenever the xlsx changes.
"""
import json
import os
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(SCRIPT_DIR, "..", "..", "docs", "KRANE-SRS-traceability.xlsx")
OUT_PATH = os.path.join(SCRIPT_DIR, "..", "..", "docs", "traceability-data.json")

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

requirements = {}
ws = wb["Requirements"]
for row in ws.iter_rows(min_row=2, values_only=True):
    flow, srs_id, requirement = row[0], row[1], row[2]
    if not srs_id:
        continue
    requirements[srs_id] = {"flow": flow, "requirement": requirement}

traceability = []
ws = wb["Traceability Report"]
for row in ws.iter_rows(min_row=4, max_row=26, max_col=9, values_only=True):
    (user_story, srs_id, test_source, test_name, test_description,
     iec_81001, expected_result, result, notes) = row
    if not srs_id:
        continue
    traceability.append({
        "userStory": user_story,
        "srsId": srs_id,
        "testSource": test_source,
        "testName": test_name,
        "testDescription": test_description,
        "iec81001": iec_81001,
        "expectedResult": expected_result,
        "lastKnownResult": result,
        "notes": notes,
        "requirement": requirements.get(srs_id, {}).get("requirement"),
        "flow": requirements.get(srs_id, {}).get("flow"),
    })

with open(OUT_PATH, "w", encoding="utf-8") as f:
    json.dump({"requirements": requirements, "traceability": traceability}, f, indent=2, ensure_ascii=False)

print(f"Wrote {len(traceability)} traceability rows to {OUT_PATH}")
